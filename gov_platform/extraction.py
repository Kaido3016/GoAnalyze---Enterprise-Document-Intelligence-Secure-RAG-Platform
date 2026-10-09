"""Bounded text extraction for text-native and scanned office documents.

PDF text extraction uses pypdf first; pages with little selectable text are
rendered with PDFium and OCRed with Tesseract. Tesseract must be installed in
the runtime image. Extraction limits are deliberate safeguards against
decompression bombs and pathological documents.
"""
from __future__ import annotations

import io
import logging
import zipfile

logger = logging.getLogger(__name__)
MAX_PDF_PAGES = 250
MAX_EXTRACTED_CHARS = 2_000_000
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_DOCX_MEMBERS = 2000
MAX_OCR_RENDER_PIXELS = 25_000_000
MIN_SELECTABLE_CHARS_PER_PAGE = 24


class ExtractionUnavailable(RuntimeError):
    """The required parser/OCR dependency is missing or could not run."""


def extract_document_text(data: bytes, content_type: str, filename: str = "") -> str:
    if not data:
        return ""
    media_type = content_type.split(";", 1)[0].strip().lower()
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if media_type == "application/pdf" or suffix == "pdf":
        return _extract_pdf(data)
    if media_type in {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"} or suffix == "docx":
        return _extract_docx(data)
    if media_type.startswith("text/") or suffix in {"txt", "csv", "md", "json", "xml", "html"}:
        return data.decode("utf-8-sig", errors="replace")[:MAX_EXTRACTED_CHARS]
    if media_type.startswith("image/") or suffix in {"png", "jpg", "jpeg", "tif", "tiff", "bmp"}:
        return _ocr_image(data)
    raise ValueError(f"unsupported_document_type:{media_type or suffix or 'unknown'}")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionUnavailable("pypdf is not installed") from exc

    reader = PdfReader(io.BytesIO(data), strict=False)
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(f"pdf_page_limit_exceeded:{MAX_PDF_PAGES}")
    pages: list[str] = []
    for page_number, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if len(text) < MIN_SELECTABLE_CHARS_PER_PAGE:
            try:
                text = _ocr_pdf_page(data, page_number).strip() or text
            except ExtractionUnavailable:
                if not text:
                    raise
                logger.warning("OCR unavailable; retaining selectable PDF text on page %s", page_number + 1)
        pages.append(f"[Page {page_number + 1}]\n{text}" if text else f"[Page {page_number + 1}]")
        if sum(len(part) for part in pages) > MAX_EXTRACTED_CHARS:
            break
    return "\n\n".join(pages)[:MAX_EXTRACTED_CHARS]


def _ocr_pdf_page(data: bytes, page_number: int) -> str:
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError as exc:
        raise ExtractionUnavailable("pypdfium2 and pytesseract are required for scanned PDFs") from exc
    try:
        pdf = pdfium.PdfDocument(data)
        page = pdf[page_number]
        width, height = page.get_size()
        if width * 1.8 * height * 1.8 > MAX_OCR_RENDER_PIXELS:
            raise ExtractionUnavailable("pdf_page_render_pixel_limit_exceeded")
        bitmap = page.render(scale=1.8, rotation=0)
        image = bitmap.to_pil()
        return pytesseract.image_to_string(image, lang="eng", timeout=45)
    except ExtractionUnavailable:
        raise
    except Exception as exc:
        if exc.__class__.__name__ == "TesseractNotFoundError":
            raise ExtractionUnavailable("Tesseract OCR executable is not installed") from exc
        raise ExtractionUnavailable(f"PDF page OCR failed: {type(exc).__name__}") from exc


def _ocr_image(data: bytes) -> str:
    try:
        from PIL import Image
        import pytesseract
    except ImportError as exc:
        raise ExtractionUnavailable("Pillow and pytesseract are required for image OCR") from exc
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.thumbnail((5000, 5000))
            return pytesseract.image_to_string(image, lang="eng", timeout=45)[:MAX_EXTRACTED_CHARS]
    except ExtractionUnavailable:
        raise
    except Exception as exc:
        if exc.__class__.__name__ == "TesseractNotFoundError":
            raise ExtractionUnavailable("Tesseract OCR executable is not installed") from exc
        raise ExtractionUnavailable(f"image OCR failed: {type(exc).__name__}") from exc


def _extract_docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_DOCX_MEMBERS:
                raise ValueError("docx_member_limit_exceeded")
            if sum(member.file_size for member in members) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ValueError("docx_uncompressed_size_limit_exceeded")
            if any(member.file_size > MAX_DOCX_UNCOMPRESSED_BYTES for member in members):
                raise ValueError("docx_member_size_limit_exceeded")
    except zipfile.BadZipFile as exc:
        raise ValueError("invalid_docx_archive") from exc
    try:
        from docx import Document
    except ImportError as exc:
        raise ExtractionUnavailable("python-docx is not installed") from exc
    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)[:MAX_EXTRACTED_CHARS]
