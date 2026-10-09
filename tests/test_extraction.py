from io import BytesIO

from docx import Document
from pypdf import PdfWriter

from gov_platform.extraction import extract_document_text


def test_plain_text_extraction_is_bounded_and_utf8():
    data = ("Environmental assessment — Québec\n" * 100).encode("utf-8")
    text = extract_document_text(data, "text/plain", "report.txt")
    assert "Québec" in text
    assert len(text) <= 2_000_000


def test_docx_paragraphs_and_tables_are_extracted():
    document = Document()
    document.add_paragraph("Applicant: Example Company")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Site"
    table.cell(0, 1).text = "North"
    output = BytesIO()
    document.save(output)
    text = extract_document_text(
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application.docx",
    )
    assert "Example Company" in text
    assert "Site | North" in text


def test_scanned_pdf_falls_back_to_ocr(monkeypatch):
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    output = BytesIO()
    writer.write(output)
    monkeypatch.setattr("gov_platform.extraction._ocr_pdf_page", lambda _data, _page: "Scanned permit text")
    text = extract_document_text(output.getvalue(), "application/pdf", "scan.pdf")
    assert "[Page 1]" in text
    assert "Scanned permit text" in text


def test_unsupported_file_type_is_rejected():
    try:
        extract_document_text(b"not a document", "application/x-executable", "payload.bin")
    except ValueError as exc:
        assert "unsupported_document_type" in str(exc)
    else:
        raise AssertionError("unsupported file type should be rejected")
